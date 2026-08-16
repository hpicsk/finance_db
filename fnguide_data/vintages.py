"""Pull-vintage manifest for the DataGuide exports in ``raw/``.

The exports are not one snapshot. Each sheet was built in its own DataGuide
session — months apart across files, minutes to hours apart *inside* one file —
and each carries its own end date and its own ticker universe. Joining two of
them silently truncates the panel to the earlier end date and can manufacture
all-NaN columns for names that listed in between. Nothing in the data announces
this: the parse succeeds and the panel is simply short.

The vintage is a property of the file, so the honest fix is not to re-pull
everything into one session — that alignment survives exactly until the next
single-file refresh, and re-anchors every back-adjusted level on the way — but
to make the vintage *readable* without paying to parse 6.5 GB of xlsx:

    python -m fnguide_data.vintages           # re-scan raw/ and rewrite vintages.csv

``vintages.csv`` is committed and the whole scan takes ~30 s, because openpyxl
streams the sheet XML and stops after the header. Consumers read the CSV:

    from fnguide_data.vintages import load, end_date
    end_date('Price data.xlsx')               # -> Timestamp('2026-08-12')
    min(end_date(f) for f in ('data0203.xlsx', 'Price data.xlsx'))

``check()`` compares the manifest against the files on disk and is what
``test_assertions.py`` runs, so a re-pull cannot land without the manifest
being regenerated in the same commit.
"""
from __future__ import annotations

import re
from itertools import islice
from pathlib import Path

import openpyxl
import pandas as pd

RAW_DIR = Path(__file__).resolve().parent / 'raw'
MANIFEST_PATH = Path(__file__).resolve().parent / 'vintages.csv'

# Header rows are found by their column-A label rather than by index. The
# English-locale exports open on a genuinely blank row where the Korean-locale
# ones carry ``Refresh``, and the readers disagree about it:
# ``CalamineSheet.to_python()`` trims that blank row, while openpyxl,
# ``CalamineSheet.iter_rows()`` and ``pd.read_excel`` keep it. Under
# ``to_python()`` the two locales therefore sit one row apart and a fixed index
# quietly returns 코드명 where 코드 was meant. Labels cost nothing and are immune
# to that and to a future layout change alike.
_REFRESH = ('Refresh',)
_TERM = ('Term', '기간')
_CODE = ('Symbol', '코드')
_END_STAMP = re.compile(r'(?:Current|최근일자)\((\d{8})\)')


def _scan_sheet(rows: list[list], sheet: str) -> dict:
    """Pull the vintage fields out of one sheet's header, whichever locale it is."""
    labelled = {str(r[0]).strip(): r for r in rows[:15] if r and r[0] is not None}

    def row_of(labels):
        for lab in labels:
            if lab in labelled:
                return labelled[lab]
        return []

    def cell(row: list, col: int) -> str:
        if len(row) <= col or row[col] is None:
            return ''
        return str(row[col]).strip()

    term = row_of(_TERM)
    end = _END_STAMP.search(cell(term, 2))
    codes = row_of(_CODE)[1:]
    return {
        'sheet': sheet,
        'built': cell(row_of(_REFRESH), 1).replace('Last Updated:', '').strip(),
        # calamine hands back the numeric start date as a float ('19990101.0').
        'term_start': cell(term, 1).removesuffix('.0'),
        'term_end': end.group(1) if end else '',
        'n_tickers': sum(1 for c in codes if c),
    }


def build(raw_dir: Path = RAW_DIR, out_path: Path = MANIFEST_PATH) -> pd.DataFrame:
    """Scan every xlsx in ``raw/`` and rewrite the manifest."""
    files = sorted(p for p in raw_dir.glob('*.xlsx') if not p.name.startswith('~$'))
    if not files:
        raise FileNotFoundError(f"no xlsx under {raw_dir}")

    records = []
    for path in files:
        # openpyxl in read-only mode, not calamine, is the exception in this
        # package: calamine parses a whole sheet before handing back a cell, so
        # reading 15 header rows out of data0204.xlsx costs ~45 s a sheet and the
        # scan runs over an hour. openpyxl streams the sheet XML and stops, which
        # takes the same read to well under a second.
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        size = path.stat().st_size
        try:
            for sheet in wb.sheetnames:
                rows = [list(r) for r in
                        islice(wb[sheet].iter_rows(max_row=15, values_only=True), 15)]
                rec = _scan_sheet(rows, sheet)
                if not rec['n_tickers']:       # an empty spare tab the export left behind
                    print(f"[vintages] {path.name} :: {sheet} -> empty, skipped")
                    continue
                rec['file'] = path.name
                rec['bytes'] = size
                records.append(rec)
                print(f"[vintages] {path.name} :: {sheet} -> {rec['term_end'] or '?'} "
                      f"({rec['n_tickers']} tickers)")
        finally:
            wb.close()

    man = pd.DataFrame(records)[['file', 'sheet', 'built', 'term_start',
                                 'term_end', 'n_tickers', 'bytes']]
    man.to_csv(out_path, index=False)
    print(f"[vintages] wrote {out_path} — {len(man)} sheets across "
          f"{man.file.nunique()} files")
    return man


def load(path: Path = MANIFEST_PATH) -> pd.DataFrame:
    """The committed manifest, with ``term_end`` parsed to a Timestamp."""
    man = pd.read_csv(path, dtype={'term_start': str, 'term_end': str})
    man['term_end'] = pd.to_datetime(man['term_end'], format='%Y%m%d', errors='coerce')
    return man


def end_date(file: str, sheet: str | None = None, path: Path = MANIFEST_PATH):
    """Latest session a file (or one of its sheets) was exported through.

    Without ``sheet`` this returns the *earliest* end date across the file's
    sheets, because that is what a panel built from the whole file is dated by —
    sheets inside one workbook do not share an end date.
    """
    man = load(path)
    rows = man[man['file'] == file]
    if sheet is not None:
        rows = rows[rows['sheet'] == sheet]
    if rows.empty:
        raise KeyError(f"{file}{f' :: {sheet}' if sheet else ''} not in {path.name}")
    return rows['term_end'].min()


def check(raw_dir: Path = RAW_DIR, path: Path = MANIFEST_PATH) -> list[str]:
    """Names of files whose presence or size disagrees with the manifest.

    Size stands in for content: a re-pull of any of these exports changes it by
    a session's worth of cells at minimum. It is a staleness signal, not a hash —
    what it guarantees is that a re-pull cannot go unnoticed, not that an
    unchanged size proves an unchanged file.
    """
    man = load(path)
    on_disk = {p.name: p.stat().st_size
               for p in raw_dir.glob('*.xlsx') if not p.name.startswith('~$')}
    recorded = man.groupby('file')['bytes'].first().to_dict()
    return sorted(set(on_disk) ^ set(recorded)
                  | {f for f in set(on_disk) & set(recorded)
                     if on_disk[f] != recorded[f]})


if __name__ == '__main__':
    build()
